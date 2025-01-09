"""
Requirements:
* Each Question Placed Above an Input Box
* Input Box remembering session
* Save to Doc option at the end
"""


import streamlit as st
from docx import Document
from io import BytesIO

# Function to create a Word document
def create_word_doc(answers):
    doc = Document()
    doc.add_heading("Future Authoring Program", level=1)
    for i, (question, answer) in enumerate(answers.items(), start=1):
        doc.add_heading(f"Question {i}", level=2)
        doc.add_paragraph(question)
        doc.add_heading("Answer", level=3)
        doc.add_paragraph(answer)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Streamlit App
st.title("Future Authoring")
st.subheader("In this exercise you will begin to create a version, in writing, of your ideal future.")
st.write("William James, the great American psychologist, once remarked that he did not know what he thought until he had written his thoughts down. When he didn’t know what to write, he wrote about anything that came to mind. Eventually, his ideas became focused and clarified. Brainstorm. Write whatever comes to mind. Don’t worry too much about sentence construction, spelling, or grammar. There will be plenty of time to write polished sentences later. Avoid criticizing what you write. Premature criticism interferes with the creative process.")

# Adding different text levels
st.header("Part 1: Imagining Your Ideal Future")

# Question 1
st.subheader("1.1 One Thing You Could Do Better")
question_1 = "If you could choose only one thing that you could do better, what would it be? Think and write for at least two minutes, then move on."
answer_1 = st.text_input(question_1)



# Question 2
st.subheader("1.2. Things to Learn About")
question_2 = "What would you like to learn more about, in the next six months? Two years? Five years? Think and write for at least two minutes, then move on."
answer_2 = st.text_input(question_2)


# Question 3
st.subheader("1.3. Improve Your Habits")
question_3 = "What habits would you like to improve? -At school? -At work? -With friends and family? -For your health? -With regards to smoking/alcohol/drug use? Think and write for at least two minutes, then move on."
answer_3 = st.text_input(question_3)

# Question 4
st.subheader("1.4. Your Social Life in the Future")
question_4 = """Friends and associates are an important part of a meaningful, productive life. Take a moment to consider your social network. 
Think about the friends you might want to have, and the connections you might want to make. It is perfectly reasonable to choose friends and associates who are good for you. 
Describe your ideal social life. Think and write for at least two minutes, then move on."""
answer_4 = st.text_input(question_4)

# Question 5
st.subheader("1.5. Your Leisure Activity in the Future")
question_5 = """Take a moment to consider the activities you would like to pursue outside of obligations such as work, family and school. The activities you choose should be worthwhile and personally meaningful. Without a plan, people often default to whatever is easiest, such as television watching, and waste their private time. If you waste 4 hours a day, which is not uncommon, then you are wasting 1400 hours a year. That is equivalent to 35 40-hour work weeks, which is almost as much as the typical individual spends at his or her job every year. 

If your time is worth \$25 per hour, then you are wasting time worth \$35,000 per year. Over a 50-year period, that is $1.8 million dollars, not counting interest or any increase in the value of your time as you develop. """
answer_5 = st.text_input(question_5)

# Question 6
st.subheader("1.6. Your Family Life in the Future")
question_6 = """Take a moment to consider your home and family life. Peaceful, harmonious family life provides people with a sense of belonging, support for their ambitions, and reciprocal purpose. 
Describe what your ideal family would be like. You can write about your parents and siblings, or about your plans for your own partner, or about your children, if any – or about all of these. What kind of partner would be good for you? 

How could you improve your relationship with your parents or siblings? Think and write for at least two minutes, then move on.
"""
answer_6 = st.text_input(question_6)

# Question 7
st.subheader("1.7. Your Career in the Future")
question_7 = """Much of what people find engaging in life is related to their careers. A good career provides security, status, interest, and the possibility of contributing to the community. Take a moment to consider your school or work careers, or both. 

Where do you want to be in six months? Two years? Five years? Why? What are you trying to accomplish? Think and write for at least two minutes, then move on.
"""
answer_7 = st.text_input(question_7)

# Question 8
st.subheader("1.8. Qualities You Admire")
question_8 = """People you automatically admire have qualities that you would like to possess or imitate. Identifying those qualities can help you determine who it is that you want to be. Take a moment to think about the two or three people you most admire. 

Who are they? Which qualities do they possess that you wish you had? Think and write for at least two minutes, then move on.
"""
answer_8 = st.text_input(question_8)

# Question 9 - First Big One
st.header("The Ideal Future: Complete Summary")
question_9 = """Now you have written briefly about your future, and have had some time to consider more specific issues. This step gives you the chance to integrate all the things that you have just thought and wrote about. Close your eyes. Daydream, if you can, and imagine your ideal future: -Who do you want to be? -What do you want to do? -Where do you want to end up? -Why do you want these things? -How do you plan to achieve your goals? -When will you put your plans into action? -Write about the ideal future that you have just imagined for 15 minutes. **Write continuously and try not to stop while you are writing. Don’t worry about spelling or grammar. You will have an opportunity to fix your mistakes later. Dream while you write, and don’t stop. Write at least until the 15 minutes have passed. Be ambitious. Imagine a life that you would regard as honourable, exciting, productive, creative and decent. Remember, you are writing only for yourself. Choose goals that you want to pursue for your own private reasons, not because someone else thinks that those goals are important. You don’t want to live someone else’s life. 

Include your deepest thoughts and feelings about all your personal goals.
"""
answer_9 = st.text_area(question_9, height= 150)

# Question 10 - Second Big One
st.header("A Future to Avoid: Complete Summary")
question_10 = """You have now written about the future you would like to have. Clearly defining your future can help reduce the uncertainty in your life, and reduce the amount of negative emotion that you chronically experience, in consequence. This is good for your confidence and for your health. Having well-defined goals also increases your chances of experiencing positive emotion, as people experience most of their hope and joy and curiosity and engagement as a consequence of pursuing valued goals (and not, as people generally think, by attaining them). It can also be very useful to deeply imagine the future you would like to avoid. You probably know people who have made very bad decisions, and who end up with a life that nobody would want. You also likely have weaknesses yourself. If you let those get out of control, then you might also end up with a miserable, painful life. Most people know how their life could go downhill if they let it. Spend some time, now, thinking about what your life would be like if you failed to define or pursue your goals, if you let your bad habits get out of control, and if you ended up miserable, resentful and bitter. Imagine your life three to five years down the road, if you failed to stay on the path you know you should be on. Use your imagination. Draw on your knowledge of the anxiety and pain you have experienced in the path, when you have betrayed yourself. Think about the people you know who have made bad decisions or remained indecisive, or who chronically deceive themselves or other people, or who let cynicism and anger dominate their lives. Where do you not want to be? Dream while you write, and don’t stop. Write at least until the 15 minutes have passed. Let yourself form a very clear picture of the undesirable future.

"""
answer_10 = st.text_area(question_10, height= 150)

# Part Two - Goal Setting

st.title("Stage 2: Specific Goal Identification: Introduction")
st.write("""In this stage, you will first be asked to define and personally title your overall future plan. Then, you will be asked to take your general plans for the ideal future and break them up into more specific goals. Each of these separate goals will also be given its own title. This step will help you clarify your goals.

Please specify a title and brief description for your ideal future as a whole. This can be as simple as “My Ideal Future,” in both fields, or, if you have something more personal in mind, you can specify that. Imagine that you are both specifying and summarizing your ambitions with this title. This will help you remember what you are aiming for. In later screens you can define, prioritize, and analyse specific goals.
Please break down your ideal future into 8 goals. You can re-word, re-write and organize the relevant material from Step 1 for your goal summaries, if you wish, or you can rely on your memory. The exercise allows you to specify a minimum of 6 goals, but people who identify 8 have better results with this exercise. These specific goals can be from a number of different domains. -A personal goal might be “I would like to be healthier.” -A career goal might be “I would like to be more interested in my job” -A social goal might be “I would like to meet more people”. The summaries you write about each goal should be reasonably brief and memorable. Make sure that each goal summary includes nothing but the most important information. You will have 10-15 minutes for this part of the exercise. Feel free to revise and edit.
""")

# Question 11-18 

question_11 = "Goal 1 title"
answer_11 = st.text_input(question_11)

question_12 = "Goal 2 title"
answer_12 = st.text_input(question_12)

question_13 = "Goal 3 title"
answer_13 = st.text_input(question_13)

question_14 = "Goal 4 title"
answer_14 = st.text_input(question_14)

question_15 = "Goal 5 title"
answer_15 = st.text_input(question_15)

question_16 = "Goal 6 title"
answer_16 = st.text_input(question_16)

question_17 = "Goal 7 title"
answer_17 = st.text_input(question_17)

question_18 = "Goal 8 title"
answer_18 = st.text_input(question_18)




# Collecting all responses
answers = {
    question_1: answer_1,
    question_2: answer_2,
    question_3: answer_3,
    question_4: answer_4,
    question_5: answer_5,
    question_6: answer_6,
    question_7: answer_7,
    question_8: answer_8,
    question_9: answer_9,
    question_10: answer_10,
    question_11: answer_11,
    question_12: answer_12,
    question_13: answer_13,
    question_14: answer_14,
    question_15: answer_15,
    question_16: answer_16,
    question_17: answer_17,
    question_18: answer_18,
}

# Button to download responses as a Word file
if st.button("Download as Word File"):
    if any(answers.values()):  # Ensure at least one answer is filled
        word_file = create_word_doc(answers)
        st.download_button(
            label="Download Word File",
            data=word_file,
            file_name="responses.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    else:
        st.warning("Please answer at least one question before downloading.")